# src/ingestion/loaders.py
import os
import hashlib
from pypdf import PdfReader
from docx import Document as DocxDocument
import pandas as pd

from src.schemas import Document


def _make_doc_id(filepath: str) -> str:
    with open(filepath, "rb") as f:
        content_hash = hashlib.md5(f.read()).hexdigest()[:12]
    return f"{os.path.basename(filepath)}_{content_hash}"


def load_pdf(filepath: str) -> Document:
    reader = PdfReader(filepath)
    # PDFs store text per page internally, so we extract and join per-page text.
    full_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return Document(
        doc_id=_make_doc_id(filepath),
        filename=os.path.basename(filepath),
        file_type="pdf",
        raw_text=full_text,
        metadata={"num_pages": len(reader.pages)},
    )


def load_docx(filepath: str) -> Document:
    doc = DocxDocument(filepath)
    # DOCX stores content as a sequence of paragraph objects — join their text.
    full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return Document(
        doc_id=_make_doc_id(filepath),
        filename=os.path.basename(filepath),
        file_type="docx",
        raw_text=full_text,
        metadata={"num_paragraphs": len(doc.paragraphs)},
    )


def load_csv(filepath: str) -> Document:
    df = pd.read_csv(filepath)
    lines = [", ".join(f"{col}: {row[col]}" for col in df.columns) for _, row in df.iterrows()]
    full_text = "\n".join(lines)
    return Document(
        doc_id=_make_doc_id(filepath),
        filename=os.path.basename(filepath),
        file_type="csv",
        raw_text=full_text,
        metadata={"num_rows": len(df), "columns": ", ".join(df.columns)},  # ← joined into a string
    )


def load_txt(filepath: str) -> Document:
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        full_text = f.read()
    return Document(
        doc_id=_make_doc_id(filepath),
        filename=os.path.basename(filepath),
        file_type="txt",
        raw_text=full_text,
        metadata={},
    )


# Dispatch table: maps file extension -> the right loader function.
# This "strategy pattern" avoids a long if/elif chain and makes adding a new format a one-line change.
LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".csv": load_csv,
    ".txt": load_txt,
}


def load_document(filepath: str) -> Document:
    ext = os.path.splitext(filepath)[1].lower()
    if ext not in LOADERS:
        raise ValueError(f"Unsupported file type: {ext}")
    return LOADERS[ext](filepath)