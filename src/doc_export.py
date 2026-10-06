import io
from docx import Document as DocxWriter


def _new_doc(title: str) -> DocxWriter:
    doc = DocxWriter()
    doc.add_heading(title, level=0)
    return doc


def _save(doc: DocxWriter) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_user_stories_docx(data: dict) -> bytes:
    doc = _new_doc("User Stories")
    for us in data.get("user_stories", []):
        p = doc.add_paragraph()
        p.add_run(f"As a {us.get('role', '')}, ").bold = True
        p.add_run(f"I want to {us.get('goal', '')} so that {us.get('benefit', '')}.")
    return _save(doc)


def build_risk_register_docx(data: dict) -> bytes:
    doc = _new_doc("Risk Register")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(["ID", "Description", "Severity", "Impact", "Mitigation Suggestion"]):
        table.rows[0].cells[i].text = h
    for r in data.get("risk_register", []):
        row = table.add_row().cells
        row[0].text = str(r.get("risk_id", ""))
        row[1].text = str(r.get("description", ""))
        row[2].text = str(r.get("severity", ""))
        row[3].text = str(r.get("impact", ""))
        row[4].text = str(r.get("mitigation_suggestion", ""))
    return _save(doc)


def build_action_items_docx(data: dict) -> bytes:
    doc = _new_doc("Action Item List")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(["Item", "Owner", "Status", "Source"]):
        table.rows[0].cells[i].text = h
    for a in data.get("action_item_list", []):
        row = table.add_row().cells
        row[0].text = str(a.get("item", ""))
        row[1].text = str(a.get("owner", ""))
        row[2].text = str(a.get("status", ""))
        row[3].text = str(a.get("source", ""))
    return _save(doc)


def build_combined_docx(user_stories: dict, risk_register: dict, action_items: dict) -> bytes:
    doc = _new_doc("Project Documentation")

    doc.add_heading("User Stories", level=1)
    for us in user_stories.get("user_stories", []):
        p = doc.add_paragraph()
        p.add_run(f"As a {us.get('role', '')}, ").bold = True
        p.add_run(f"I want to {us.get('goal', '')} so that {us.get('benefit', '')}.")

    doc.add_heading("Risk Register", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(["ID", "Description", "Severity", "Impact", "Mitigation Suggestion"]):
        table.rows[0].cells[i].text = h
    for r in risk_register.get("risk_register", []):
        row = table.add_row().cells
        row[0].text = str(r.get("risk_id", ""))
        row[1].text = str(r.get("description", ""))
        row[2].text = str(r.get("severity", ""))
        row[3].text = str(r.get("impact", ""))
        row[4].text = str(r.get("mitigation_suggestion", ""))

    doc.add_heading("Action Item List", level=1)
    table2 = doc.add_table(rows=1, cols=4)
    table2.style = "Light Grid Accent 1"
    for i, h in enumerate(["Item", "Owner", "Status", "Source"]):
        table2.rows[0].cells[i].text = h
    for a in action_items.get("action_item_list", []):
        row = table2.add_row().cells
        row[0].text = str(a.get("item", ""))
        row[1].text = str(a.get("owner", ""))
        row[2].text = str(a.get("status", ""))
        row[3].text = str(a.get("source", ""))

    return _save(doc)